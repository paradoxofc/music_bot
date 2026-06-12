import io
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

REPORT_HEADERS = (
    'ID пользователя',
    'Имя',
    'Username',
    'Язык',
    'Telegram Premium',
    'Статус',
    'Дата подписки',
    'Ref',
    'Скачано треков',
)

COLUMN_WIDTHS = (14, 22, 18, 10, 16, 18, 20, 16, 16)


def _format_value(value) -> str:
    if value is None or value == '':
        return '—'
    return str(value)


def _format_name(row: dict) -> str:
    parts = [row.get('first_name'), row.get('last_name')]
    name = ' '.join(p for p in parts if p)
    return name or '—'


def _format_username(username: str | None) -> str:
    if not username:
        return '—'
    return f'@{username}'


def _format_premium(is_premium: bool | None) -> str:
    if is_premium is True:
        return 'Да'
    if is_premium is False:
        return 'Нет'
    return '—'


def _format_status(is_active: bool | None) -> str:
    if is_active is False:
        return 'Заблокировал бота'
    return 'Подписан'


def _format_datetime(dt, tz: ZoneInfo | None = None) -> str:
    if not dt:
        return '—'
    if tz:
        dt = dt.astimezone(tz)
    return dt.strftime('%d.%m.%Y %H:%M')


def _user_to_row(row: dict, tz: ZoneInfo | None = None) -> tuple:
    return (
        row.get('chat_id'),
        _format_name(row),
        _format_username(row.get('username')),
        _format_value(row.get('language')),
        _format_premium(row.get('is_premium')),
        _format_status(row.get('is_active')),
        _format_datetime(row.get('created_at'), tz),
        _format_value(row.get('source')),
        row.get('downloads_count', 0),
    )


def build_users_report_xlsx(users: list[dict], tz_name: str = 'Europe/Moscow') -> io.BytesIO:
    """Собирает Excel-отчёт по списку пользователей."""
    tz = ZoneInfo(tz_name)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Пользователи'

    header_font = Font(bold=True)
    for col, header in enumerate(REPORT_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = COLUMN_WIDTHS[col - 1]

    for row_idx, user in enumerate(users, start=2):
        for col_idx, value in enumerate(_user_to_row(user, tz), start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical='center')

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(REPORT_HEADERS))}{max(1, len(users) + 1)}'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def report_filename() -> str:
    stamp = datetime.now(ZoneInfo('Europe/Moscow')).strftime('%Y-%m-%d_%H-%M')
    return f'users_report_{stamp}.xlsx'
